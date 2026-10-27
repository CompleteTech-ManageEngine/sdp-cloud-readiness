"""Scope of Work placeholders, Markdown subset and branded outputs; synthetic data only."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from sdp_audit.branded import write_branded_pack
from sdp_audit.branding import Brand
from sdp_audit.environment_analysis import analyze_environment
from sdp_audit.report import render_report_documents
from sdp_audit.scope_of_work import INVENTORY, Figures, body_html, document_title, parse, render_html
from sdp_audit.synthetic import demo_environment

SOURCE = """# Fictional Migration

## Scope of Work

**Prepared for {{client}}**
**Prepared by {{consultant}}**

{{consultant}} shall assess {{portal_count}} portals at {{source_url}} as of {{assessment_date}}.

{{inventory_table}}

- Templates: {{inventory_total:incident_templates+service_templates}}
- See [V1](https://help.servicedeskplus.com/op-to-od-migration) and <script>x</script>

<!-- pagebreak -->

| Party | Responsibility |
| --- | --- |
| {{consultant}} | Lead delivery |
"""


class ScopeOfWorkTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.documents = render_report_documents(analyze_environment(demo_environment("multi")))
        cls.data = json.loads(cls.documents["assessment.json"])
        # Contract figures need complete inventories; the demo deliberately includes an inaccessible portal.
        keys = {key for key, _ in INVENTORY}
        cls.data["portals"] = [p for p in cls.data["portals"]
                               if all(r.get("list_complete") is True for r in p["coverage"] if r.get("key") in keys)]
        cls.documents = {**cls.documents, "assessment.json": json.dumps(cls.data)}
        cls.brand = Brand(name="Fictional Advisory", website="fictional.example", prepared_for="Client <One>")

    def test_placeholders_fill_from_brand_and_evidence(self):
        figures = Figures(self.data, self.brand)
        text = figures.fill(SOURCE)
        self.assertNotIn("{{", text)
        self.assertIn("**Prepared for Client <One>**", text)
        self.assertIn("Fictional Advisory shall assess", text)
        self.assertIn("| **Total** |", text)
        self.assertEqual(figures.used["{{consultant}}"], "Fictional Advisory")

    def test_unknown_placeholder_missing_client_and_absent_impact_stop_rendering(self):
        with self.assertRaises(ValueError):
            Figures(self.data, self.brand).fill("{{not_a_figure}}")
        with self.assertRaises(ValueError):
            Figures(self.data, Brand(name="X")).fill("{{client}}")
        with self.assertRaises(ValueError):
            Figures(self.data, self.brand).fill("{{impact_total:NO_SUCH_RULE}}")

    def test_incomplete_inventory_is_refused(self):
        data = json.loads(json.dumps(self.data))
        data["portals"][0]["coverage"] = [dict(row, list_complete=False) for row in data["portals"][0]["coverage"]]
        with self.assertRaises(ValueError):
            Figures(data, self.brand).fill("{{inventory_total:reports}}")

    def test_markdown_subset_escapes_and_breaks(self):
        blocks = parse(Figures(self.data, self.brand).fill(SOURCE))
        kinds = [kind for kind, _ in blocks]
        self.assertEqual(document_title(blocks), "Fictional Migration")
        self.assertIn("pagebreak", kinds)
        self.assertEqual(kinds.count("table"), 2)
        html = body_html(blocks)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)
        self.assertIn('<a href="https://help.servicedeskplus.com/op-to-od-migration">V1</a>', html)
        self.assertIn("Client &lt;One&gt;</strong><br>", html)
        page = render_html(blocks, self.brand)
        self.assertIn("script-src 'none'", page)
        self.assertIn('class="brand-band"', page)

    def test_pack_renders_scope_of_work_when_present(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            for name, text in self.documents.items():
                (folder / name).write_text(text, encoding="utf-8")
            (folder / "scope-of-work.md").write_text(SOURCE, encoding="utf-8")
            written = write_branded_pack(folder, self.brand, pdf=False, log=lambda _: None)
            self.assertIn("Client-One-ServiceDesk-Plus-Cloud-Scope-of-Work.html", written)
            manifest = json.loads((folder / "branded/Fictional-Advisory/branded-manifest.json").read_text(encoding="utf-8"))
            self.assertIn("scope-of-work.md", [item["name"] for item in manifest["sources"]])
            self.assertEqual(manifest["scope_of_work_figures"]["{{client}}"], "Client <One>")
            if importlib.util.find_spec("docx"):
                import docx
                document = docx.Document(str(written["Client-One-ServiceDesk-Plus-Cloud-Scope-of-Work.docx"]))
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn("Fictional Advisory shall assess", text)
                self.assertEqual(document.core_properties.author, "Fictional Advisory")
            else:
                self.assertTrue(any("python-docx" in note for note in manifest["skipped"]))

    def test_pack_without_scope_of_work_is_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            for name, text in self.documents.items():
                (folder / name).write_text(text, encoding="utf-8")
            written = write_branded_pack(folder, self.brand, pdf=False, log=lambda _: None)
            self.assertFalse(any("Scope-of-Work" in name for name in written))
