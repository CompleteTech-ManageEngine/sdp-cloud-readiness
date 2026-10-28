"""PERT effort estimate: evidence drivers, arithmetic, validation and the editable workbook."""

from __future__ import annotations

import importlib.util
import json
import math
from pathlib import Path
import tempfile
import unittest

from sdp_audit.branded import write_branded_pack
from sdp_audit.branding import Brand
from sdp_audit.effort_estimate import Z80, load_estimate, markdown
from sdp_audit.environment_analysis import analyze_environment
from sdp_audit.report import render_report_documents
from sdp_audit.scope_of_work import INVENTORY
from sdp_audit.synthetic import demo_environment

ESTIMATE = {
    "title": "Fictional Effort Estimate",
    "contingency_percent": 10,
    "hourly_rate": 100,
    "deliverables": {"D01": "Mobilization", "D02": "Design"},
    "tasks": [
        {"deliverable": "D01", "task": "Plan", "driver": "fixed", "optimistic": 6, "most_likely": 9, "pessimistic": 18},
        {"deliverable": "D02", "task": "Per portal", "driver": "portal_count", "optimistic": 1, "most_likely": 2, "pessimistic": 3},
        {"deliverable": "D02", "task": "Reports", "driver": "inventory_total:reports", "optimistic": 0.5, "most_likely": 1, "pessimistic": 3},
        {"deliverable": "D02", "task": "Manual", "driver": 4, "optimistic": 1, "most_likely": 1, "pessimistic": 1, "basis": "Not in evidence"},
    ],
}


class EffortEstimateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.documents = render_report_documents(analyze_environment(demo_environment("multi")))
        data = json.loads(cls.documents["assessment.json"])
        keys = {key for key, _ in INVENTORY}
        data["portals"] = [p for p in data["portals"]
                           if all(r.get("list_complete") is True for r in p["coverage"] if r.get("key") in keys)]
        cls.data = data
        cls.documents = {**cls.documents, "assessment.json": json.dumps(data)}
        cls.brand = Brand(name="Fictional Advisory", prepared_for="Client")

    def load(self, estimate):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "effort-estimate.json"
            path.write_text(json.dumps(estimate), encoding="utf-8")
            return load_estimate(path, self.data, self.brand)

    def test_pert_arithmetic_and_drivers(self):
        estimate = self.load(ESTIMATE)
        portals = len(self.data["portals"])
        reports = sum(next(r for r in p["coverage"] if r["key"] == "reports")["records_observed"] for p in self.data["portals"])
        quantities = [t.quantity for t in estimate.tasks]
        self.assertEqual(quantities, [1, portals, reports, 4])
        expected = 10 + portals * 2 + reports * (0.5 + 4 + 3) / 6 + 4
        deviation = math.sqrt(2 ** 2 + (portals * 2 / 6) ** 2 + (reports * 2.5 / 6) ** 2)
        totals = estimate.totals()
        self.assertAlmostEqual(totals["expected"], expected)
        self.assertAlmostEqual(totals["deviation"], deviation)
        self.assertAlmostEqual(totals["p80"], expected + Z80 * deviation)
        self.assertAlmostEqual(totals["with_contingency"], expected * 1.1)
        text = markdown(estimate, self.brand, self.data)
        self.assertIn(f"**{totals['with_contingency']:,.0f} hours**", text)
        self.assertIn(f"| **Total estimate** | **{totals['with_contingency']:,.0f}** |", text)
        self.assertIn("Indicative cost", text)
        self.assertIn(f"| Per portal | {portals} portals |", text)
        self.assertIn("| Manual | 4 (allowance) | 4 | 4\u2060\u2013\u20604 | Not in evidence |", text)

    def test_client_document_hides_working_detail(self):
        text = markdown(self.load(ESTIMATE), self.brand, self.data)
        for internal in ("h/unit", "Std. deviation", "σ", "PERT)", "driver", "inventory_total", "Draft", "variance"):
            self.assertNotIn(internal, text)
        self.assertIn("PERT (three-point) method", text)
        self.assertIn("## Estimate by deliverable", text)

    def test_validation(self):
        for change in ({"tasks": []}, {"deliverables": {"X1": "bad"}},
                       {"tasks": [dict(ESTIMATE["tasks"][0], deliverable="D09")]},
                       {"tasks": [dict(ESTIMATE["tasks"][0], optimistic=10, most_likely=5)]},
                       {"tasks": [dict(ESTIMATE["tasks"][0], pessimistic=-1)]},
                       {"tasks": [dict(ESTIMATE["tasks"][0], driver="unknown_driver")]},
                       {"tasks": [dict(ESTIMATE["tasks"][0], driver="impact_total:NO_SUCH_RULE")]},
                       {"contingency_percent": True}):
            with self.subTest(change=change):
                with self.assertRaises(ValueError):
                    self.load({**ESTIMATE, **change})

    def test_pack_writes_estimate_in_every_format(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            for name, text in self.documents.items():
                (folder / name).write_text(text, encoding="utf-8")
            (folder / "effort-estimate.json").write_text(json.dumps(ESTIMATE), encoding="utf-8")
            written = write_branded_pack(folder, self.brand, pdf=False, log=lambda _: None)
            self.assertIn("Client-Effort-Estimate.html", written)
            self.assertIn("Client-Effort-Estimate.md", written)
            manifest = json.loads((folder / "branded/Fictional-Advisory/branded-manifest.json").read_text(encoding="utf-8"))
            self.assertIn("effort-estimate.json", [item["name"] for item in manifest["sources"]])
            self.assertGreater(manifest["effort_estimate_hours"]["expected"], 0)
            if importlib.util.find_spec("openpyxl"):
                import openpyxl
                book = openpyxl.load_workbook(written["Client-Effort-Estimate.xlsx"])
                self.assertEqual(book.sheetnames, ["Summary", "Tasks", "Method"])
                tasks = book["Tasks"]
                header = next(r for r in range(1, 30) if tasks.cell(row=r, column=1).value == "Deliverable")
                first = header + 1
                self.assertEqual(tasks.cell(row=first, column=12).value,
                                 f'=IF(OR($D{first}="",$E{first}="",$F{first}="",$G{first}=""),"",$D{first}*$H{first})')
                self.assertEqual(tasks.cell(row=first, column=6).fill.start_color.rgb[-6:], "FFF2CC")
                formulas = [c.value for row in book["Summary"].iter_rows() for c in row if isinstance(c.value, str)]
                self.assertTrue(any(f.startswith("=SUMIF(Tasks!$A$") for f in formulas))
                self.assertTrue(any("0.8416" in f for f in formulas))
                self.assertTrue(book.calculation.fullCalcOnLoad)

    def test_displayed_hours_add_up(self):
        import re
        from sdp_audit.effort_estimate import rounded_parts
        self.assertEqual(rounded_parts([141.6, 221.5, 110.4], 474), [142, 222, 110])
        self.assertEqual(sum(rounded_parts([0.4, 0.4, 0.4, 0.4], 2)), 2)
        text = markdown(self.load(ESTIMATE), self.brand, self.data)
        number = lambda value: int(value.replace(",", ""))  # noqa: E731
        parts = [number(h) for h in re.findall(r"^\| D\d\d [^|]+\| ([\d,]+) \|", text, re.M)]
        total = number(re.search(r"^\| \*\*Total\*\* \| \*\*([\d,]+)\*\*", text, re.M).group(1))
        self.assertEqual(sum(parts), total)
        for section in re.split(r"^### ", text, flags=re.M)[1:]:
            rows = re.findall(r"^\| (?!Task |---|\*\*Subtotal)[^|]+\|[^|]+\| ([\d,]+) \|", section, re.M)
            subtotal = number(re.search(r"\*\*Subtotal\*\* \| \| \*\*([\d,]+)\*\*", section).group(1))
            self.assertEqual(sum(number(r) for r in rows), subtotal)

    def test_detailed_document_keeps_full_working_table(self):
        from sdp_audit.effort_estimate import detailed_markdown
        estimate = self.load(ESTIMATE)
        detailed = detailed_markdown(estimate, self.brand, self.data)
        client = markdown(estimate, self.brand, self.data)
        for column in ("O h/unit", "M h/unit", "P h/unit", "Std. dev. h", "Quantity basis", "Std. deviation"):
            self.assertIn(column, detailed)
            self.assertNotIn(column, client)
        self.assertIn("| Reports | Query reports |", detailed)
        self.assertIn(f"**{estimate.totals()['expected']:,.1f} hours** (PERT)", detailed)
        self.assertIn("detailed working estimate", detailed)

    def test_pack_writes_client_and_detailed_estimates(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            for name, text in self.documents.items():
                (folder / name).write_text(text, encoding="utf-8")
            (folder / "effort-estimate.json").write_text(json.dumps(ESTIMATE), encoding="utf-8")
            written = write_branded_pack(folder, self.brand, pdf=False, log=lambda _: None)
            for stem in ("Client-Effort-Estimate", "Client-Effort-Estimate-Detailed"):
                self.assertIn(f"{stem}.html", written)
                self.assertIn(f"{stem}.md", written)
