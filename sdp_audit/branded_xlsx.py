"""Branded Excel companion to findings.csv and coverage.csv (requires openpyxl).

The CSVs remain the authoritative evidence. This writes a formatted copy: a
cover sheet with the brand band, stripe, prepared by/for and summary counts,
then every row and column of both CSVs with values unchanged. Counts become
numbers; identifiers, flags and JSON evidence stay text and are never
interpreted as formulas.
"""

from __future__ import annotations

from collections import Counter
import csv
from io import BytesIO
from pathlib import Path
import re
import zipfile

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .branding import Brand

TITLE = "Cloud Readiness Assessment: Findings & Coverage"
FONT = "Arial"
INT_RE = re.compile(r"(0|-?[1-9]\d{0,14})")
FORMULA_LEAD = ("=", "+", "-", "@")
MAX_WIDTH = 60
WIDE_COLUMNS = {"evidence": 80, "last_response": 80, "diagnostics": 70, "summary": 60, "action": 55}
MAX_ROW_HEIGHT = 120.0
SEVERITY_ORDER = ["critical", "high", "medium", "low", "info"]


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


class _Style:
    def __init__(self, brand: Brand):
        self.ink, self.ink_2, self.muted = brand.hex("ink"), brand.hex("ink_2"), brand.hex("ink_3")
        self.band, self.band_ink = brand.hex("band"), brand.hex("band_ink")
        self.accent, self.accent_2 = brand.hex("accent"), brand.hex("accent_2")
        self.light, self.line = brand.hex("surface"), brand.hex("line")
        thin = Side(style="thin", color=self.line)
        self.cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        edge = Side(style="thin", color=self.band)
        self.header_border = Border(left=edge, right=edge, top=edge, bottom=Side(style="medium", color=self.accent))

    def font(self, size: float = 10, color: str | None = None, bold: bool = False, italic: bool = False) -> Font:
        return Font(name=FONT, size=size, color=color or self.ink_2, bold=bold, italic=italic)


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return [], []
    header, body = rows[0], rows[1:]
    for number, row in enumerate(body, start=2):
        if len(row) != len(header):
            raise ValueError(f"{path.name}: line {number} has {len(row)} fields, expected {len(header)}")
        for value in row:
            if ILLEGAL_CHARACTERS_RE.search(value) or len(value) > 32767:
                raise ValueError(f"{path.name}: line {number} holds a value Excel cannot store")
    return header, body


def _numeric_columns(header: list[str], body: list[list[str]]) -> set[int]:
    result = set()
    for index, name in enumerate(header):
        if name == "id" or name.endswith("_id") or name == "key":
            continue
        values = [row[index] for row in body if row[index] != ""]
        if values and all(INT_RE.fullmatch(v) for v in values):
            result.add(index)
    return result


def _put_text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"
    if value.startswith(FORMULA_LEAD):
        cell.quotePrefix = True


def _column_width(name: str, values: list[str]) -> tuple[float, bool]:
    longest = max([len(name)] + [max((len(part) for part in v.split("\n")), default=0) for v in values])
    cap = WIDE_COLUMNS.get(name, MAX_WIDTH)
    return min(max(longest + 2, len(name) + 4, 8), cap), longest + 2 > cap


def _estimate_lines(value: str, width: float) -> int:
    chars = max(int(width * 1.1), 1)
    return sum(max(1, -(-len(part) // chars)) for part in value.split("\n"))


def _header_footer(ws, brand: Brand, s: _Style, left: str, right: str) -> None:
    ws.oddHeader.left.text = left
    ws.oddHeader.right.text = right
    ws.oddFooter.left.text = " · ".join(part for part in (brand.name, brand.website) if part)
    ws.oddFooter.right.text = "Page &P of &N"
    for part in (ws.oddHeader.left, ws.oddHeader.right, ws.oddFooter.left, ws.oddFooter.right):
        part.font, part.size, part.color = FONT, 8, s.muted


def _data_sheet(wb, s: _Style, brand: Brand, title: str, source: Path, header: list[str], body: list[list[str]]) -> None:
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = s.accent if title == "Findings" else s.accent_2
    if not header:
        ws.cell(row=1, column=1, value="The source CSV is empty.")
        return
    numbers = _numeric_columns(header, body)
    widths, wraps = [], []
    for index, name in enumerate(header):
        width, wrap = _column_width(name, [row[index] for row in body])
        widths.append(width)
        wraps.append(wrap)
        ws.column_dimensions[get_column_letter(index + 1)].width = width
    header_font, header_fill = s.font(10, s.band_ink, bold=True), _fill(s.band)
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for index, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=index)
        _put_text(cell, name)
        cell.font, cell.fill, cell.border, cell.alignment = header_font, header_fill, s.header_border, header_align
    ws.row_dimensions[1].height = 30
    body_font, light = s.font(9), _fill(s.light)
    wrap_align, plain_align = Alignment(vertical="top", wrap_text=True), Alignment(vertical="top")
    number_align = Alignment(horizontal="right", vertical="top")
    for r, row in enumerate(body, start=2):
        lines = 1
        for c, value in enumerate(row):
            cell = ws.cell(row=r, column=c + 1)
            if value == "":
                cell.value = None
            elif c in numbers:
                cell.value = int(value)
                cell.number_format = "0"
            else:
                _put_text(cell, value)
            cell.font, cell.border = body_font, s.cell_border
            if r % 2 == 1:
                cell.fill = light
            if c in numbers:
                cell.alignment = number_align
            elif wraps[c]:
                cell.alignment = wrap_align
                lines = max(lines, _estimate_lines(value, widths[c]))
            else:
                cell.alignment = plain_align
        if lines > 1:
            ws.row_dimensions[r].height = min(12.5 * lines + 2, MAX_ROW_HEIGHT)
    last = f"{get_column_letter(len(header))}{len(body) + 1}"
    ws.freeze_panes = "B2" if title == "Findings" else "A2"
    ws.auto_filter.ref = f"A1:{last}"
    ws.print_title_rows = "1:1"
    ws.print_area = f"A1:{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    prefix = (brand.prepared_for + " · ") if brand.prepared_for else ""
    _header_footer(ws, brand, s, prefix + TITLE, f"{title} ({source.name})")


def _cover_sheet(wb, s: _Style, brand: Brand, folder_name: str, findings, coverage, sources: tuple[Path, Path]) -> None:
    ws = wb.active
    ws.title = "Cover"
    ws.sheet_properties.tabColor = s.band
    ws.sheet_view.showGridLines = False
    widths = [3, 30, 14, 14, 14, 14, 14, 3]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    last_col = len(widths)
    for row, height in {1: 14, 2: 46, 3: 10, 4: 30, 5: 20, 6: 12}.items():
        ws.row_dimensions[row].height = height
        for col in range(1, last_col + 1):
            ws.cell(row=row, column=col).fill = _fill(s.band)
    if brand.logo:
        logo = XLImage(BytesIO(brand.logo))
        width, height = brand.logo_size
        logo.height = 44
        logo.width = round(44 * width / height) if height else 132
        ws.add_image(logo, "B2")
    else:
        ws.cell(row=2, column=2, value=brand.name).font = s.font(16, s.band_ink, bold=True)
    if brand.prepared_for:
        tag = ws.cell(row=2, column=7, value=brand.prepared_for.upper())
        tag.font = s.font(9, s.accent_2, bold=True)
        tag.alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=4, column=2, value=TITLE).font = s.font(20, s.band_ink, bold=True)
    ws.cell(row=4, column=2).alignment = Alignment(vertical="center")
    ws.cell(row=5, column=2, value="ServiceDesk Plus cloud readiness · companion evidence workbook").font = s.font(10, s.accent_2)
    ws.row_dimensions[7].height = 4
    for col in range(1, last_col + 1):
        ws.cell(row=7, column=col).fill = _fill(s.accent if col <= 3 else s.accent_2)

    row = 9
    details = [("Prepared for", brand.prepared_for)] if brand.prepared_for else []
    details += [
        ("Prepared by", brand.name),
        ("Contact", " · ".join(part for part in (brand.website, brand.email) if part)),
        ("Assessment folder", folder_name),
        ("Source files", f"{sources[0].name} ({len(findings[1])} rows × {len(findings[0])} columns); "
                         f"{sources[1].name} ({len(coverage[1])} rows × {len(coverage[0])} columns)"),
    ]
    for label, value in details:
        if not value:
            continue
        ws.cell(row=row, column=2, value=label).font = s.font(10, s.muted, bold=True)
        ws.cell(row=row, column=3, value=value).font = s.font(10, s.ink)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        row += 1

    def heading(text: str) -> None:
        nonlocal row
        row += 1
        ws.cell(row=row, column=2, value=text).font = s.font(13, s.ink, bold=True)
        for col in range(2, 8):
            ws.cell(row=row, column=col).border = Border(bottom=Side(style="medium", color=s.accent))
        row += 1

    def table(columns: list[str], rows: list[list]) -> None:
        nonlocal row
        for index, name in enumerate(columns):
            cell = ws.cell(row=row, column=2 + index, value=name)
            cell.font, cell.fill, cell.border = s.font(10, s.band_ink, bold=True), _fill(s.band), s.header_border
            cell.alignment = Alignment(horizontal="left" if index == 0 else "right", vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1
        for number, values in enumerate(rows):
            for index, value in enumerate(values):
                cell = ws.cell(row=row, column=2 + index, value=value)
                cell.font, cell.border = s.font(10), s.cell_border
                cell.alignment = Alignment(horizontal="left" if index == 0 else "right")
                if number % 2:
                    cell.fill = _fill(s.light)
            row += 1
        for index, value in enumerate(["Total"] + [sum(r[i] for r in rows) for i in range(1, len(columns))]):
            cell = ws.cell(row=row, column=2 + index, value=value)
            cell.font = s.font(10, s.ink, bold=True)
            cell.border = Border(top=Side(style="thin", color=s.ink), bottom=Side(style="thin", color=s.line))
            cell.alignment = Alignment(horizontal="left" if index == 0 else "right")
        row += 1

    f_header, f_body = findings
    if "status" in f_header and "severity" in f_header:
        status_i, severity_i = f_header.index("status"), f_header.index("severity")
        severities = sorted({r[severity_i] for r in f_body},
                            key=lambda v: SEVERITY_ORDER.index(v) if v in SEVERITY_ORDER else 99)
        pairs = Counter((r[status_i], r[severity_i]) for r in f_body)
        statuses = Counter(r[status_i] for r in f_body)
        heading("Findings by status")
        table(["Status"] + [v.title() or "—" for v in severities] + ["Findings"],
              [[status] + [pairs[(status, v)] for v in severities] + [count] for status, count in statuses.most_common()])
    c_header, c_body = coverage
    outcome_key = "outcome" if "outcome" in c_header else "status" if "status" in c_header else None
    complete_key = "list_complete" if "list_complete" in c_header else "list_complete_for_api_scope" if "list_complete_for_api_scope" in c_header else None
    if outcome_key and complete_key:
        outcome_i, complete_i = c_header.index(outcome_key), c_header.index(complete_key)
        outcomes = Counter(r[outcome_i] for r in c_body)
        listed = Counter((r[outcome_i], r[complete_i].lower()) for r in c_body)
        heading("Coverage rows by outcome")
        table(["Outcome", "List complete", "List incomplete", "Rows"],
              [[o, listed[(o, "true")], listed[(o, "false")], n] for o, n in outcomes.most_common()])

    heading("About this workbook")
    notes = [
        f"The CSV files in the assessment folder ({sources[0].name}, {sources[1].name}) remain the authoritative evidence; "
        "their SHA-256 hashes are recorded in report-manifest.json and branded/branded-manifest.json. "
        "This workbook is a formatted copy for reading and filtering.",
        "The Findings and Coverage sheets hold every row and column of the CSVs with values unchanged. Counts are stored as "
        "numbers; identifiers, flags and JSON evidence are stored as text. Empty CSV fields are empty cells.",
        "Very long cells (JSON evidence and API responses) are complete in the cell, but rows are capped in height; select a "
        "cell or widen the row to read all of it.",
    ]
    for text in notes:
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = s.font(9.5)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        ws.row_dimensions[row].height = 40
        row += 1
    row += 1
    ws.cell(row=row, column=2, value=brand.contact).font = s.font(8.5, s.muted)
    ws.print_area = f"A1:{get_column_letter(last_col)}{row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0.2


def _add_company(path: Path, company: str) -> None:
    """openpyxl does not write docProps/app.xml Company; add it after saving."""
    from xml.sax.saxutils import escape
    temp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/app.xml":
                text = data.decode("utf-8")
                if "<Company>" not in text:
                    text = re.sub(r"<Properties([^>]*)/>", lambda m: f"<Properties{m.group(1)}></Properties>", text)
                    text = re.sub(r"<Properties[^>]*>", lambda m: f"{m.group(0)}<Company>{escape(company)}</Company>", text, count=1)
                data = text.encode("utf-8")
            target.writestr(item, data)
    temp.replace(path)


def write_evidence_workbook(findings_path: Path, coverage_path: Path, target: Path, brand: Brand, folder_name: str) -> Path:
    findings, coverage = _read_csv(findings_path), _read_csv(coverage_path)
    style = _Style(brand)
    wb = openpyxl.Workbook()
    _cover_sheet(wb, style, brand, folder_name, findings, coverage, (findings_path, coverage_path))
    _data_sheet(wb, style, brand, "Findings", findings_path, *findings)
    _data_sheet(wb, style, brand, "Coverage", coverage_path, *coverage)
    wb.active = 0
    props = wb.properties
    props.title = " · ".join(part for part in (brand.prepared_for, TITLE) if part)
    props.subject = "ServiceDesk Plus cloud readiness assessment evidence (formatted copy of findings.csv and coverage.csv)"
    props.creator = props.lastModifiedBy = brand.name
    props.category = "Assessment evidence"
    temp = target.with_name(target.name + ".partial.xlsx")
    wb.save(temp)
    _add_company(temp, brand.name)
    temp.replace(target)
    return target
