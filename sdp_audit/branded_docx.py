"""Editable branded Word copy of the Scope of Work (requires python-docx).

Uses the same parsed blocks as the HTML/PDF rendering, so the three formats
carry identical text. Header: brand band with logo and document type; footer:
contact line and "Page n of N".
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .branding import Brand

FONT = "Arial"


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.lstrip("#").upper())


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    props.append(shading)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def _hyperlink(paragraph, url: str, text: str, colour: str) -> None:
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                   is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, attrs in (("w:color", {"w:val": colour.lstrip("#")}), ("w:u", {"w:val": "single"})):
        element = OxmlElement(tag)
        for key, value in attrs.items():
            element.set(qn(key), value)
        props.append(element)
    run.append(props)
    content = OxmlElement("w:t")
    content.text = text
    content.set(qn("xml:space"), "preserve")
    run.append(content)
    link.append(run)
    paragraph._p.append(link)


def _inline(paragraph, text: str, brand: Brand, size: float | None = None, colour: str | None = None,
            bold: bool = False) -> None:
    for part in re.split(r"(\*\*.+?\*\*|\[[^\]]+\]\(https?://[^)\s]+\))", text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", part)
        if link:
            _hyperlink(paragraph, link.group(2), link.group(1), brand.tokens["link"])
            continue
        strong = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if strong else part)
        run.bold = bold or strong
        run.font.name = FONT
        if size:
            run.font.size = Pt(size)
        if colour:
            run.font.color.rgb = _rgb(colour)


def _header(section, brand: Brand) -> None:
    t = brand.tokens
    header = section.header
    header.is_linked_to_previous = False
    table = header.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    for cell in (left, right):
        _shade(cell, t["band"])
    para = left.paragraphs[0]
    if brand.logo:
        if brand.logo_background:
            _shade(left, brand.logo_background)
        para.add_run().add_picture(BytesIO(brand.logo), height=Inches(0.42))
    else:
        _inline(para, brand.name, brand, 12, t["band_ink"], bold=True)
    doc_type = right.paragraphs[0]
    doc_type.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _inline(doc_type, "Scope of Work", brand, 14, t["band_ink"])
    if brand.prepared_for:
        ref = right.add_paragraph()
        ref.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _inline(ref, brand.prepared_for.upper(), brand, 7.5, t["accent_2"], bold=True)
    header.paragraphs[0].text = ""


def _footer(section, brand: Brand) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    _inline(para, brand.contact + "    ", brand, 7.5, brand.tokens["ink_3"])
    _inline(para, "Page ", brand, 7.5, brand.tokens["ink_3"])
    _field(para, "PAGE")
    _inline(para, " of ", brand, 7.5, brand.tokens["ink_3"])
    _field(para, "NUMPAGES")


def write_scope_docx(blocks: list[tuple[str, Any]], brand: Brand, title: str, target: Path) -> Path:
    t = brand.tokens
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.85)
    section.top_margin, section.bottom_margin = Inches(1.1), Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.font.color.rgb = _rgb(t["ink_2"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    _header(section, brand)
    _footer(section, brand)
    sizes = {1: (24, t["ink"]), 2: (16, t["ink"]), 3: (12.5, t["accent"]), 4: (11, t["ink"])}
    for kind, value in blocks:
        if kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "heading":
            level, text = value
            size, colour = sizes.get(level, (10.5, t["ink"]))
            para = document.add_paragraph()
            para.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.keep_with_next = True
            _inline(para, text, brand, size, colour, bold=True)
        elif kind == "para":
            para = document.add_paragraph()
            para.paragraph_format.space_after = Pt(6)
            for index, line in enumerate(value):
                if index:
                    para.add_run().add_break()
                _inline(para, line.strip(), brand)
        elif kind == "list":
            for item in value:
                para = document.add_paragraph(style="List Bullet")
                _inline(para, item, brand)
        elif kind == "table":
            head, *rows = value
            table = document.add_table(rows=1 + len(rows), cols=len(head))
            table.style = "Table Grid"
            for col, text in enumerate(head):
                cell = table.rows[0].cells[col]
                _shade(cell, t["band"])
                _inline(cell.paragraphs[0], text, brand, 9, t["band_ink"], bold=True)
            for r, row in enumerate(rows, start=1):
                for col in range(len(head)):
                    cell = table.rows[r].cells[col]
                    if r % 2 == 0:
                        _shade(cell, t["surface"])
                    _inline(cell.paragraphs[0], row[col] if col < len(row) else "", brand, 9)
            document.add_paragraph()
    props = document.core_properties
    props.title = title
    props.author = props.last_modified_by = brand.name
    props.subject = "ServiceDesk Plus Cloud migration Scope of Work" + (f" for {brand.prepared_for}" if brand.prepared_for else "")
    temp = target.with_name(target.name + ".partial.docx")
    document.save(str(temp))
    temp.replace(target)
    return target
